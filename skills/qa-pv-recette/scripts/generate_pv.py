"""Génère un PV de recette Word à partir d'un JSON et du modèle assets/pv_recette_template.docx.

Usage : py generate_pv.py data.json sortie.docx [--template modele.docx]
Voir ../references/example.json pour le format des données.
"""
import argparse
import copy
import json
import os
import re
import sys

import docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "..", "assets", "pv_recette_template.docx")

# ordre des cases à cocher dans le modèle
CHECKS = ["fonctionnels", "tnr", "accepte", "accepte_reserves", "refuse",
          "bloquante", "majeure", "mineure"]
TOKEN = re.compile(r"\{\{(r\.)?(\w+)\}\}")


def uniq(row):
    seen, out = [], []
    for c in row.cells:
        if c._tc not in seen:
            seen.append(c._tc)
            out.append(c)
    return out


def fill_text(root, values, row_scope=False):
    """Remplace les jetons {{x}} (ou {{r.x}} si row_scope) dans tous les w:t sous root."""
    for t in root.iter(qn("w:t")):
        if not t.text or "{{" not in t.text:
            continue

        def sub(m):
            is_row, key = bool(m.group(1)), m.group(2)
            if is_row != row_scope:
                return m.group(0)
            return str(values.get(key, "") or "")
        t.text = TOKEN.sub(sub, t.text)


def fill_rows(table, proto_idx, items, keys):
    """Clone la ligne prototype une fois par élément. Liste vide : ligne laissée vierge."""
    proto = table.rows[proto_idx]._tr
    parent = proto.getparent()
    if not items:
        fill_text(proto, {}, row_scope=True)
        return
    anchor = proto
    for it in items:
        row = copy.deepcopy(proto)
        fill_text(row, {k: it.get(k, "") for k in keys}, row_scope=True)
        anchor.addnext(row)
        anchor = row
    parent.remove(proto)


def fill_paragraphs(cell, items, prefix=""):
    """Clone le paragraphe prototype de la cellule pour chaque élément."""
    proto = cell.paragraphs[0]._p
    if not items:
        items = ["Néant"]
    anchor = proto
    for it in items:
        p = copy.deepcopy(proto)
        fill_text(p, {"item": prefix + it}, row_scope=True)
        anchor.addnext(p)
        anchor = p
    proto.getparent().remove(proto)


def tick(doc, wanted):
    boxes = list(doc.element.body.iter(qn("w:default")))
    if len(boxes) != len(CHECKS):
        sys.exit(f"modèle inattendu : {len(boxes)} cases à cocher au lieu de {len(CHECKS)}")
    for key, e in zip(CHECKS, boxes):
        e.set(qn("w:val"), "1" if key in wanted else "0")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("data")
    ap.add_argument("dst")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    a = ap.parse_args()
    data = json.load(open(a.data, encoding="utf-8"))
    d = docx.Document(a.template)
    T = d.tables

    # tableaux à lignes dynamiques (indices du modèle)
    fill_rows(T[2], 5, data.get("bugs", []), ["id", "libelle", "sprint"])      # bugs d'abord
    fill_rows(T[2], 3, data.get("stories", []), ["id", "libelle", "sprint"])   # puis stories
    fill_paragraphs(uniq(T[3].rows[1])[1], data.get("perimetre_tnr", []), "- ")
    fill_paragraphs(uniq(T[3].rows[2])[1], data.get("hors_perimetre", []))
    fill_rows(T[5], 2, data.get("reserves", []), ["id", "libelle"])
    fill_rows(T[6], 2, data.get("anomalies", []),
              ["id", "libelle", "severite", "statut", "contournement"])
    fill_rows(T[7], 2, data.get("non_validees", []),
              ["id", "libelle", "commentaire", "perimetre"])
    fill_rows(T[8], 5, data.get("reserves_decision", []), ["date", "numero", "descriptif"])

    # jetons simples : corps, en-têtes
    values = {k: v for k, v in data.items() if isinstance(v, str)}
    values.setdefault("version", "V 1.0")
    values.setdefault("sponsor_role", "Responsable métier / Sponsor")
    roots = [d.element.body]
    for s in d.sections:
        for h in (s.header, s.first_page_header, s.even_page_header, s.footer):
            roots.append(h._element)
    for r in roots:
        fill_text(r, values)

    # cases à cocher
    wanted = set(data.get("type_recette", ["fonctionnels", "tnr"]))
    wanted.add(data.get("decision", "accepte"))
    if data.get("reserves_statut"):
        wanted.add(data["reserves_statut"])
    tick(d, wanted)

    # capture du graphe de couverture : remplace l'image placeholder
    if data.get("graphe"):
        from PIL import Image
        import io
        for rel in d.part.rels.values():
            if rel.reltype.endswith("/image") and "image5" in rel.target_ref:
                buf = io.BytesIO()
                Image.open(data["graphe"]).convert("RGB").save(buf, "PNG")
                rel.target_part._blob = buf.getvalue()
                break

    # jetons non renseignés : prévenir
    left = set()
    for r in roots:
        for t in r.iter(qn("w:t")):
            left.update(TOKEN.findall(t.text or ""))
    if left:
        print("jetons restants :", sorted(k for _, k in left))

    d.core_properties.title = f"PV de recette {values.get('projet', '')}".strip()
    d.save(a.dst)
    print("ok ->", a.dst)


if __name__ == "__main__":
    main()
