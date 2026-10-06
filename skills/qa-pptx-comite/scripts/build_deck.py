"""JSON de contenu -> présentation .pptx (titre, puces, graphique, tableau).
Usage : py build_deck.py contenu.json sortie.pptx [--template template.pptx]"""
import json, argparse
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src"); ap.add_argument("dst"); ap.add_argument("--template")
    a = ap.parse_args()
    data = json.load(open(a.src, encoding="utf-8"))
    prs = Presentation(a.template) if a.template else Presentation()
    if not a.template:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    title_layout, content_layout, only_title = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5]

    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = data["title"]
    if len(s.placeholders) > 1:
        s.placeholders[1].text = data.get("subtitle", "")

    W, H = prs.slide_width, prs.slide_height
    for sd in data["slides"]:
        if sd["type"] == "bullets":
            s = prs.slides.add_slide(content_layout)
            s.shapes.title.text = sd["title"]
            tf = s.placeholders[1].text_frame
            tf.text = sd["bullets"][0]
            for b in sd["bullets"][1:]:
                tf.add_paragraph().text = b
        elif sd["type"] == "chart":
            s = prs.slides.add_slide(only_title)
            s.shapes.title.text = sd["title"]
            cd = CategoryChartData(); cd.categories = sd["categories"]
            cd.add_series(sd.get("series", "Tests"), sd["values"])
            ch = s.shapes.add_chart(XL_CHART_TYPE.PIE if sd.get("pie") else XL_CHART_TYPE.COLUMN_CLUSTERED,
                                    Inches(1), Inches(1.6), W - Inches(2), H - Inches(2.2), cd).chart
            ch.has_legend = bool(sd.get("pie"))
            if ch.has_legend:
                ch.legend.position = XL_LEGEND_POSITION.RIGHT
            ch.plots[0].has_data_labels = True
        elif sd["type"] == "table":
            s = prs.slides.add_slide(only_title)
            s.shapes.title.text = sd["title"]
            rows = [sd["header"]] + sd["rows"]
            t = s.shapes.add_table(len(rows), len(sd["header"]), Inches(0.8), Inches(1.6),
                                   W - Inches(1.6), Inches(0.4) * len(rows)).table
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    cell = t.cell(i, j); cell.text = str(v)
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(14); run.font.bold = (i == 0)
    prs.save(a.dst)
    print(f"{len(prs.slides)} slides -> {a.dst}")

if __name__ == "__main__":
    main()
