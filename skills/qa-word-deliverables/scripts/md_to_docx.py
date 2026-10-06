"""Convertit un Markdown simple (titres, listes, tableaux, **gras**) en .docx.
Usage : py md_to_docx.py entree.md sortie.docx [--template template.docx]"""
import re, sys, argparse
from docx import Document

def add_runs(par, text):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            par.add_run(part).bold = (i % 2 == 1)

def flush_table(doc, rows):
    rows = [r for r in rows if not re.fullmatch(r"\|?[\s:\-|]+\|?", r)]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    if not cells:
        return
    t = doc.add_table(rows=len(cells), cols=max(len(r) for r in cells))
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    for i, r in enumerate(cells):
        for j, c in enumerate(r):
            p = t.cell(i, j).paragraphs[0]
            add_runs(p, c)
            if i == 0:
                for run in p.runs:
                    run.bold = True

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src"); ap.add_argument("dst"); ap.add_argument("--template")
    a = ap.parse_args()
    doc = Document(a.template) if a.template else Document()
    table = []
    for line in open(a.src, encoding="utf-8").read().splitlines():
        if line.strip().startswith("|"):
            table.append(line); continue
        if table:
            flush_table(doc, table); table = []
        m = re.match(r"(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"\s*[-*]\s+(\[.\]\s+)?", line):
            add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"\s*[-*]\s+", "", line, count=1))
        elif re.match(r"\s*\d+\.\s+", line):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"\s*\d+\.\s+", "", line, count=1))
        elif line.strip():
            add_runs(doc.add_paragraph(), line)
    if table:
        flush_table(doc, table)
    doc.save(a.dst)
    print(f"OK -> {a.dst}")

if __name__ == "__main__":
    main()
