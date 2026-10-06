"""Rapport JUnit (Playwright) -> classeur Excel Détail + Synthèse.
Usage : py junit_to_xlsx.py results/junit.xml synthese_campagne.xlsx"""
import sys, re, xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule

def group_of(name):
    m = re.search(r"api[\\/]+([^\\/]+)[\\/]", name)
    return m.group(1) if m else (name.split(" ")[0] or "autre")

def main(src, dst):
    root = ET.parse(src).getroot()
    rows = []
    for tc in root.iter("testcase"):
        cls = tc.get("classname", "")
        status, detail = "OK", ""
        if tc.find("skipped") is not None:
            status, detail = "SKIPPED", tc.find("skipped").get("message", "")
        elif tc.find("failure") is not None or tc.find("error") is not None:
            el = tc.find("failure") if tc.find("failure") is not None else tc.find("error")
            status, detail = "KO", (el.get("message") or "")[:300]
        rows.append([group_of(cls), cls, tc.get("name", ""), status, float(tc.get("time", 0)), detail])

    wb = Workbook()
    ws = wb.active; ws.title = "Détail"
    ws.append(["Groupe", "Fichier", "Test", "Statut", "Durée (s)", "Détail"])
    for r in rows: ws.append(r)
    for c in ws[1]: c.font = Font(bold=True)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col, w in zip("ABCDEF", [16, 40, 60, 10, 10, 80]): ws.column_dimensions[col].width = w
    rng = f"D2:D{len(rows)+1}"
    for val, color in [('"OK"', "C6EFCE"), ('"KO"', "FFC7CE"), ('"SKIPPED"', "D9D9D9")]:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[val],
                                      fill=PatternFill("solid", fgColor=color)))

    s = wb.create_sheet("Synthèse")
    s.append(["Groupe", "Total", "OK", "KO", "Skipped", "Taux OK / exécutés"])
    groups = sorted({r[0] for r in rows})
    n = len(rows) + 1
    for i, g in enumerate(groups, start=2):
        s.append([g, f'=COUNTIFS(Détail!A2:A{n},A{i})',
                  f'=COUNTIFS(Détail!A2:A{n},A{i},Détail!D2:D{n},"OK")',
                  f'=COUNTIFS(Détail!A2:A{n},A{i},Détail!D2:D{n},"KO")',
                  f'=COUNTIFS(Détail!A2:A{n},A{i},Détail!D2:D{n},"SKIPPED")',
                  f'=IFERROR(C{i}/(C{i}+D{i}),0)'])
    last = len(groups) + 1
    s.append(["TOTAL"] + [f"=SUM({c}2:{c}{last})" for c in "BCDE"] + [f"=IFERROR(C{last+1}/(C{last+1}+D{last+1}),0)"])
    for c in s[1] + s[last + 1]: c.font = Font(bold=True)
    for row in s.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in row: c.number_format = "0.0%"
    s.column_dimensions["A"].width = 22
    wb.save(dst)
    print(f"{len(rows)} tests -> {dst}")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
